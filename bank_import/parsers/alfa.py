from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from bank_import.models import BankOperation
from .common import booked_at, money, text


RAW_FIELDS = [
    "Дата", "Номер документа", "Дебет", "Кредит", "Контрагент", "ИНН", "КПП",
    "Счёт", "БИК", "Наименование банка", "Назначение платежа", "Код дебитора", "Тип документа",
]


def parse(path: str | Path) -> list[BankOperation]:
    source = Path(path)
    workbook = load_workbook(source, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    if "выписка по сч" not in text(sheet["A1"].value).casefold():
        workbook.close()
        raise ValueError("Альфа-Банк: не найден заголовок формы в A1")
    own_account = text(sheet["B1"].value)
    operations: list[BankOperation] = []
    for row_number, values in enumerate(sheet.iter_rows(min_row=13, max_col=13, values_only=True), start=13):
        values = list(values)
        if not any(text(value) for value in values):
            continue
        debit, credit = money(values[2]), money(values[3])
        if (debit > 0) == (credit > 0):
            raise ValueError(f"Альфа-Банк, строка {row_number}: должен быть заполнен ровно один из Дебет/Кредит")
        direction = "debit" if debit > 0 else "credit"
        operations.append(BankOperation(
            bank="alfa",
            source_path=str(source),
            source_sheet=sheet.title,
            source_row=row_number,
            account=own_account,
            booked_at=booked_at(values[0]),
            document_number=text(values[1]),
            direction=direction,
            amount=debit if direction == "debit" else credit,
            currency="RUB",
            description=text(values[12]),
            purpose=text(values[10]),
            counterparty_name=text(values[4]),
            counterparty_inn=text(values[5]),
            counterparty_account=text(values[7]),
            raw={name: text(value) for name, value in zip(RAW_FIELDS, values)},
        ))
    workbook.close()
    return operations


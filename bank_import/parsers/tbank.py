from __future__ import annotations

import csv
from pathlib import Path
from typing import Iterable

from openpyxl import load_workbook

from bank_import.models import BankOperation
from .common import booked_at, money, normalize_currency, text


HEADERS = [
    "Номер счёта", "Тип операции", "Дата проведения", "Номер документа", "Валюта операции",
    "Сумма в валюте счёта", "Валюта счёта", "Описание операции", "Назначение платежа",
    "Счет плательщика", "ИНН плательщика", "КПП плательщика", "Наименование плательщика",
    "БИК банка плательщика", "Корр. счет плательщика", "Счет получателя", "Договор получателя",
    "ИНН получателя", "КПП получателя", "Наименование получателя", "БИК банка получателя",
    "Корр. счет получателя", "Счет контрагента", "ИНН контрагента", "Наименование контрагента",
    "БИК банка контрагента", "Дебет", "Кредит",
]


def _operation(path: Path, sheet: str, row_number: int, row: dict[str, object]) -> BankOperation:
    missing = [header for header in HEADERS if header not in row]
    if missing:
        raise ValueError(f"Т-Банк: отсутствуют поля: {', '.join(missing)}")
    debit = money(row["Дебет"])
    credit = money(row["Кредит"])
    if (debit > 0) == (credit > 0):
        raise ValueError(f"Т-Банк, строка {row_number}: должен быть заполнен ровно один из Дебет/Кредит")
    direction = "debit" if debit > 0 else "credit"
    amount = debit if direction == "debit" else credit
    raw = {header: text(row.get(header)) for header in HEADERS}
    return BankOperation(
        bank="tbank",
        source_path=str(path),
        source_sheet=sheet,
        source_row=row_number,
        account=text(row["Номер счёта"]),
        booked_at=booked_at(row["Дата проведения"]),
        document_number=text(row["Номер документа"]),
        direction=direction,
        amount=amount,
        currency=normalize_currency(row["Валюта счёта"]),
        description=text(row["Описание операции"]),
        purpose=text(row["Назначение платежа"]),
        counterparty_name=text(row["Наименование контрагента"]),
        counterparty_inn=text(row["ИНН контрагента"]),
        counterparty_account=text(row["Счет контрагента"]),
        raw=raw,
    )


def _csv_rows(path: Path) -> tuple[str, Iterable[tuple[int, dict[str, object]]]]:
    handle = path.open("r", encoding="utf-8-sig", newline="")
    reader = csv.DictReader(handle)
    if reader.fieldnames != HEADERS:
        handle.close()
        raise ValueError("Т-Банк CSV: порядок или состав 28 полей не соответствует форме")

    def iterator() -> Iterable[tuple[int, dict[str, object]]]:
        try:
            for number, row in enumerate(reader, start=2):
                if any(text(value) for value in row.values()):
                    yield number, row
        finally:
            handle.close()

    return "CSV", iterator()


def _xlsx_rows(path: Path) -> tuple[str, list[tuple[int, dict[str, object]]]]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    header_row = None
    for row_number, values in enumerate(sheet.iter_rows(min_row=1, max_row=30, values_only=True), start=1):
        normalized = [text(value) for value in values[: len(HEADERS)]]
        if normalized == HEADERS:
            header_row = row_number
            break
    if header_row is None:
        workbook.close()
        raise ValueError("Т-Банк XLSX: строка с 28 заголовками не найдена")
    rows: list[tuple[int, dict[str, object]]] = []
    for row_number, values in enumerate(sheet.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
        values = list(values[: len(HEADERS)])
        if not any(text(value) for value in values):
            continue
        rows.append((row_number, dict(zip(HEADERS, values))))
    sheet_name = sheet.title
    workbook.close()
    return sheet_name, rows


def parse(path: str | Path) -> list[BankOperation]:
    source = Path(path)
    sheet, rows = _csv_rows(source) if source.suffix.lower() == ".csv" else _xlsx_rows(source)
    return [_operation(source, sheet, number, row) for number, row in rows]


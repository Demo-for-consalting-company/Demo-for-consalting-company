from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from bank_import.models import BankOperation
from .common import booked_at, money, split_party_block, text


RAW_FIELDS = [
    "Дата проводки", "Счет дебета", "Счет кредита", "Сумма по дебету", "Сумма по кредиту",
    "Номер документа", "ВО", "Банк", "Назначение платежа",
]


def parse(path: str | Path) -> list[BankOperation]:
    source = Path(path)
    workbook = load_workbook(source, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    title = " ".join(text(sheet.cell(row=row, column=column).value) for row in range(1, 9) for column in range(1, 22))
    if "выписка операций по лицевому счету" not in title.casefold():
        workbook.close()
        raise ValueError("СберБизнес: не найден заголовок формы")
    own_account = text(sheet["M5"].value)
    operations: list[BankOperation] = []
    for row_number, values in enumerate(sheet.iter_rows(min_row=12, max_col=21, values_only=True), start=12):
        values = list(values)
        date_value = values[1]
        if text(date_value).casefold() == "б/с":
            break
        debit, credit = money(values[9]), money(values[13])
        if debit == 0 and credit == 0:
            continue
        if (debit > 0) == (credit > 0):
            raise ValueError(f"СберБизнес, строка {row_number}: должен быть заполнен ровно один денежный поток")
        direction = "debit" if debit > 0 else "credit"
        counterparty_block = values[8] if direction == "debit" else values[4]
        counterparty_account, counterparty_inn, counterparty_name = split_party_block(counterparty_block)
        raw_values = [values[1], values[4], values[8], values[9], values[13], values[14], values[16], values[17], values[20]]
        operations.append(BankOperation(
            bank="sber",
            source_path=str(source),
            source_sheet=sheet.title,
            source_row=row_number,
            account=own_account,
            booked_at=booked_at(date_value),
            document_number=text(values[14]),
            direction=direction,
            amount=debit if direction == "debit" else credit,
            currency="RUB",
            description=text(values[16]),
            purpose=text(values[20]),
            counterparty_name=counterparty_name,
            counterparty_inn=counterparty_inn,
            counterparty_account=counterparty_account,
            raw={name: text(value) for name, value in zip(RAW_FIELDS, raw_values)},
        ))
    workbook.close()
    return operations


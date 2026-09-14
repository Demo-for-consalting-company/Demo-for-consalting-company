from __future__ import annotations

from pathlib import Path

from openpyxl import load_workbook

from .models import BankOperation
from .parsers import alfa, sber, tbank
from .parsers.common import text


def detect_bank(path: str | Path) -> str:
    source = Path(path)
    if source.suffix.lower() == ".csv":
        return "tbank"
    if source.suffix.lower() not in {".xlsx", ".xlsm"}:
        raise ValueError(f"Неподдерживаемый формат файла: {source.suffix}")
    workbook = load_workbook(source, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    rows = list(sheet.iter_rows(min_row=1, max_row=20, max_col=30, values_only=True))
    workbook.close()
    normalized_rows = [[text(value).casefold() for value in row] for row in rows]
    flat = " ".join(value for row in normalized_rows for value in row)
    if "выписка операций по лицевому счету" in flat:
        return "sber"
    for row in normalized_rows:
        if "номер счёта" in row and "тип операции" in row and "дата проведения" in row:
            return "tbank"
    if "выписка по счёту" in flat or "выписка по счету" in flat:
        return "alfa"
    raise ValueError("Не удалось определить форму: ожидается Т-Банк, Альфа-Банк или СберБизнес")


def parse_statement(path: str | Path) -> list[BankOperation]:
    bank = detect_bank(path)
    parser = {"tbank": tbank.parse, "alfa": alfa.parse, "sber": sber.parse}[bank]
    return parser(path)

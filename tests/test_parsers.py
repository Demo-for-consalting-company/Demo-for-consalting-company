from __future__ import annotations

import csv
import json
import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook

from bank_import.detect import detect_bank, parse_statement
from bank_import.parsers.tbank import HEADERS
from bank_import.pipeline import JOURNAL_COLUMNS, process_files


class ParserTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.root = Path(self.temp_dir.name)

    def tearDown(self) -> None:
        self.temp_dir.cleanup()

    def _tbank(self, duplicate: bool = False) -> Path:
        path = self.root / "tbank.csv"
        row = {header: "" for header in HEADERS}
        row.update({
            "Номер счёта": "DEMO-ACCOUNT",
            "Тип операции": "Кредит",
            "Дата проведения": "01.08.2026",
            "Номер документа": "T-1",
            "Валюта операции": "643",
            "Сумма в валюте счёта": "1000",
            "Валюта счёта": "643",
            "Описание операции": "Оплата клиента",
            "Назначение платежа": "Оплата по проекту Орион",
            "Счет контрагента": "DEMO-COUNTER",
            "ИНН контрагента": "DEMO-INN",
            "Наименование контрагента": "ООО Тестовый клиент",
            "Кредит": "1000",
            "Дебет": "0",
        })
        with path.open("w", encoding="utf-8-sig", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=HEADERS)
            writer.writeheader()
            writer.writerow(row)
            if duplicate:
                writer.writerow(row)
        return path

    def _alfa(self) -> Path:
        path = self.root / "alfa.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Выписка по счёту"
        sheet["A1"], sheet["B1"] = "Выписка по счёту", "DEMO-ALFA"
        sheet["A11"] = "Дата"
        sheet.append([])
        values = ["01.08.2026", "A-1", 500, None, "ООО Подрядчик", "DEMO-INN", "DEMO-KPP", "DEMO-ACC", "DEMO-BIK", "Тестовый банк", "Оплата подрядчику по проекту Орион", None, "Платежное поручение"]
        for column, value in enumerate(values, start=1):
            sheet.cell(13, column, value)
        workbook.save(path)
        return path

    def _sber(self) -> Path:
        path = self.root / "sber.xlsx"
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Выписка"
        sheet["B5"] = "ВЫПИСКА ОПЕРАЦИЙ ПО ЛИЦЕВОМУ СЧЕТУ"
        sheet["M5"] = "DEMO-SBER"
        values = [None, "01.08.2026 10:15", None, None, "DEMO-SBER\nDEMO-INN-OWN\nООО Тестовая компания", None, None, None, "DEMO-COUNTER\nDEMO-INN\nООО Арендодатель", 750, None, None, None, None, "S-1", None, "01", "DEMO-BIK Тестовый банк", None, None, "Аренда офиса"]
        for column, value in enumerate(values, start=1):
            sheet.cell(12, column, value)
        sheet["B13"] = "б/с"
        workbook.save(path)
        return path

    def test_detects_and_parses_all_three_forms(self) -> None:
        paths = [self._tbank(), self._alfa(), self._sber()]
        self.assertEqual([detect_bank(path) for path in paths], ["tbank", "alfa", "sber"])
        parsed = [parse_statement(path)[0] for path in paths]
        self.assertEqual([item.direction for item in parsed], ["credit", "debit", "debit"])
        self.assertEqual([str(item.amount) for item in parsed], ["1000.00", "500.00", "750.00"])
        self.assertTrue(all(item.operation_id for item in parsed))

    def test_pipeline_uses_project_mapping_and_removes_duplicate(self) -> None:
        input_path = self._tbank(duplicate=True)
        mapping_path = self.root / "mapping.json"
        mapping_path.write_text(json.dumps({
            "project_id": "test",
            "rules": [{
                "name": "Орион",
                "when": {"direction": "credit", "contains_any": ["орион"]},
                "set": {"article": "Оплата клиента", "project": "Орион", "cash_flow_line": "Покупатели", "pnl_line": "Выручка", "balance_account": "Покупатели"},
            }],
        }, ensure_ascii=False), encoding="utf-8")
        result = process_files([input_path], mapping_path)
        self.assertEqual(result.input_count, 2)
        self.assertEqual(len(result.ready), 1)
        self.assertEqual(len(result.duplicates), 1)
        self.assertEqual(result.ready[0].mapping.project, "Орион")
        record = result.ready[0].journal_record()
        self.assertEqual(set(record), set(JOURNAL_COLUMNS))
        self.assertEqual(record["Счёт денежных средств"], "Основной счёт")
        self.assertEqual(record["Балансовая статья"], "Покупатели")
        self.assertNotIn(input_path.name, record["ID источника"])


if __name__ == "__main__":
    unittest.main()
